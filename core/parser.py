"""文档解析模块（阶段2实现）。

统一解析三类输入：
  A类：有文字层的 PDF —— 用 PyMuPDF 按坐标提取文本块，手写分栏/阅读顺序还原。
  B类：无文字层的扫描 PDF —— 渲染为图片后用 PaddleOCR 的版面检测(LayoutDetection)
       + 文本检测识别(PaddleOCR) 独立模型组合，再用 paddlex 内置的 XY-Cut 算法
       （sort_by_xycut）还原阅读顺序（双栏场景已正确处理）。
  C类：Word 文档 —— 用 python-docx 按段落顺序读取。

只做解析，不分块、不接校对LLM（对应阶段3/4）。
"""

from __future__ import annotations

import re
import statistics
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import docx
import fitz  # PyMuPDF
import numpy as np
from PIL import Image

import config


class UnsupportedFormatError(Exception):
    """不支持的文件格式。"""


class NoTextLayerError(Exception):
    """文档无文字层，且调用方要求禁用OCR（ocr='off'）。"""


@dataclass
class ParsedBlock:
    page: int
    block_index: int
    text: str
    block_type: str  # 'paragraph'/'heading'/'footnote'/'table'/'figure_caption'/'other'
    source_location: str
    ocr_confidence: float | None = None


@dataclass
class ParsedDocument:
    file_name: str
    file_type: str  # 'pdf'/'docx'
    total_pages: int
    blocks: list[ParsedBlock] = field(default_factory=list)
    layout_mode: str = "single"  # 'single'/'double'/'mixed'
    text_source: str = "native"  # 'native'/'ocr'/'mixed'
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# 对外入口
# ---------------------------------------------------------------------------

def parse_document(
    file_path: str | Path,
    force_layout: str = "auto",
    spread_order: str = "normal",
    ocr: str = "auto",
) -> ParsedDocument:
    """解析文档，返回带页码/位置信息、按正确阅读顺序排列的结构化文本。"""
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pdf":
        return _parse_pdf(path, force_layout=force_layout, spread_order=spread_order, ocr=ocr)
    raise UnsupportedFormatError(f"不支持的文件格式：{suffix}")


# ---------------------------------------------------------------------------
# C类：Word
# ---------------------------------------------------------------------------

def _parse_docx(path: Path) -> ParsedDocument:
    document = docx.Document(str(path))
    blocks: list[ParsedBlock] = []
    idx = 0
    for i, para in enumerate(document.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style is not None else ""
        block_type = "heading" if (style_name.startswith("Heading") or style_name == "Title") else "paragraph"
        blocks.append(
            ParsedBlock(
                page=0,
                block_index=idx,
                text=text,
                block_type=block_type,
                source_location=f"第{i}段",
                ocr_confidence=None,
            )
        )
        idx += 1
    return ParsedDocument(
        file_name=path.name,
        file_type="docx",
        total_pages=len(document.paragraphs),
        blocks=blocks,
        layout_mode="single",
        text_source="native",
        warnings=[],
    )


# ---------------------------------------------------------------------------
# PDF 总入口：逐页判断文字层，分派到 A类/B类通道
# ---------------------------------------------------------------------------

def _parse_pdf(path: Path, force_layout: str, spread_order: str, ocr: str) -> ParsedDocument:
    doc = fitz.open(str(path))
    warnings: list[str] = []
    logical_pages: list[dict] = []  # {"source","page_no","mode","blocks"}
    native_pending: list[tuple[dict, list[dict], float]] = []  # (占位, 原始块, 页宽)

    logical_page_no = 0
    try:
        for phys_index in range(doc.page_count):
            page = doc[phys_index]
            char_count = len(page.get_text().strip())
            page_is_native = char_count >= config.TEXT_LAYER_MIN_CHARS

            if ocr == "off" and not page_is_native:
                raise NoTextLayerError(f"第{phys_index + 1}页无文字层，且 ocr='off' 不允许OCR")

            if ocr == "force":
                use_ocr = True
            elif ocr == "off":
                use_ocr = False
            else:
                use_ocr = not page_is_native

            if not use_ocr:
                logical_page_no += 1
                raw_blocks, width = _extract_native_page_raw(page)
                placeholder: dict = {"source": "native", "page_no": logical_page_no, "mode": None, "blocks": None}
                logical_pages.append(placeholder)
                native_pending.append((placeholder, raw_blocks, width))
            else:
                img = _render_page_image(page, config.OCR_RENDER_DPI)
                sub_images = _maybe_split_spread(img, spread_order, warnings, phys_index + 1)
                for sub_img in sub_images:
                    logical_page_no += 1
                    raw_blocks, page_conf, mode = _ocr_and_order(sub_img, force_layout)
                    if page_conf is not None and page_conf < config.OCR_PAGE_LOW_CONFIDENCE_THRESHOLD:
                        warnings.append(f"第{logical_page_no}页整体OCR识别置信度偏低({page_conf:.2f})")
                    finalized = [
                        {
                            "text": b["text"],
                            "block_type": b["block_type"],
                            "source_location": _source_location(logical_page_no, mode, b.get("column")),
                            "confidence": b["confidence"],
                        }
                        for b in raw_blocks
                    ]
                    logical_pages.append({"source": "ocr", "page_no": logical_page_no, "mode": mode, "blocks": finalized})
    finally:
        doc.close()

    if native_pending:
        stripped_list = _strip_headers_footers([rb for _, rb, _ in native_pending])
        for (placeholder, _, width), stripped in zip(native_pending, stripped_list):
            finalized, mode = _finalize_native_page(stripped, width, placeholder["page_no"], force_layout)
            placeholder["mode"] = mode
            placeholder["blocks"] = finalized

    all_blocks: list[ParsedBlock] = []
    idx = 0
    text_sources: set[str] = set()
    layout_modes: list[str] = []
    for lp in logical_pages:
        text_sources.add(lp["source"])
        layout_modes.append(lp["mode"])
        for b in lp["blocks"]:
            all_blocks.append(
                ParsedBlock(
                    page=lp["page_no"],
                    block_index=idx,
                    text=b["text"],
                    block_type=b["block_type"],
                    source_location=b["source_location"],
                    ocr_confidence=b["confidence"],
                )
            )
            idx += 1

    text_source = text_sources.pop() if len(text_sources) == 1 else "mixed"
    layout_mode = _majority_layout_mode(layout_modes)

    return ParsedDocument(
        file_name=path.name,
        file_type="pdf",
        total_pages=logical_page_no,
        blocks=all_blocks,
        layout_mode=layout_mode,
        text_source=text_source,
        warnings=warnings,
    )


def _majority_layout_mode(page_modes: list[str]) -> str:
    """按多数页面的分栏结果汇总文档级 layout_mode。

    单页分栏判断是启发式的，个别页（如整页大表格/图片、正文块很少的页）
    偶尔会判断有误；只要多数页一致就采用多数结果，避免一两页误判就把
    全文档标成 'mixed'。只有多数页本身就分歧较大时才真正视为 'mixed'。
    """
    if not page_modes:
        return "single"
    counts = Counter(page_modes)
    mode, count = counts.most_common(1)[0]
    return mode if count / len(page_modes) > 0.5 else "mixed"


def _source_location(page_no: int, mode: str, column: str | None) -> str:
    if mode != "double" or column is None:
        return f"第{page_no}页"
    if column == "left":
        return f"第{page_no}页左栏"
    if column == "right":
        return f"第{page_no}页右栏"
    return f"第{page_no}页通栏"


# ---------------------------------------------------------------------------
# 通用：在给定区间内寻找一条“空白分栏带”（按坐标区间覆盖率判断，A类原生坐标
# 与 B类版面块坐标共用同一套逻辑）
# ---------------------------------------------------------------------------

_GAP_RESOLUTION = 200


def _find_extent_gap(
    extents: list[tuple[float, float]],
    total_width: float,
    band: tuple[float, float],
    min_gap_ratio: float,
) -> float | None:
    """在 extents（[(x0,x1), ...]）里找页面中部空白竖带，返回分栏线 x 坐标。"""
    if not extents or total_width <= 0:
        return None
    bin_width = total_width / _GAP_RESOLUTION
    covered = [False] * _GAP_RESOLUTION
    for x0, x1 in extents:
        i0 = max(0, min(_GAP_RESOLUTION - 1, int(x0 / bin_width)))
        i1 = max(0, min(_GAP_RESOLUTION - 1, int(x1 / bin_width)))
        for i in range(i0, i1 + 1):
            covered[i] = True

    band_i0 = max(0, int(band[0] * _GAP_RESOLUTION))
    band_i1 = min(_GAP_RESOLUTION, int(band[1] * _GAP_RESOLUTION))

    best: tuple[int, int] | None = None
    start: int | None = None
    for i in range(band_i0, band_i1):
        if not covered[i]:
            if start is None:
                start = i
        elif start is not None:
            if best is None or (i - start) > (best[1] - best[0]):
                best = (start, i)
            start = None
    if start is not None:
        if best is None or (band_i1 - start) > (best[1] - best[0]):
            best = (start, band_i1)

    if best is None:
        return None
    gap_width = (best[1] - best[0]) * bin_width
    if gap_width / total_width < min_gap_ratio:
        return None
    return (best[0] + best[1]) / 2 * bin_width


# ---------------------------------------------------------------------------
# A类：有文字层 PDF
# ---------------------------------------------------------------------------

_NUMERIC_ZONE_RE = re.compile(r"^[\dIVXLCDMivxlcdm\-\.\s]{1,10}$")


def _zone_of(bbox: tuple[float, float, float, float], height: float) -> str:
    y0, y1 = bbox[1], bbox[3]
    if y1 <= height * config.HEADER_FOOTER_ZONE_RATIO:
        return "top"
    if y0 >= height * (1 - config.HEADER_FOOTER_ZONE_RATIO):
        return "bottom"
    return "body"


def _extract_native_page_raw(page: "fitz.Page") -> tuple[list[dict], float]:
    data = page.get_text("dict")
    width, height = data["width"], data["height"]
    out: list[dict] = []
    for b in data["blocks"]:
        if b.get("type") != 0:
            continue  # 跳过图片块（过滤图片区域）
        lines_text = []
        sizes = []
        for line in b["lines"]:
            t = "".join(s["text"] for s in line["spans"])
            if t.strip():
                lines_text.append(t.strip())
            sizes.extend(s["size"] for s in line["spans"])
        text = " ".join(lines_text)
        if not text:
            continue
        avg_size = sum(sizes) / len(sizes) if sizes else 0.0
        bbox = tuple(b["bbox"])
        out.append({"text": text, "bbox": bbox, "avg_size": avg_size, "zone": _zone_of(bbox, height)})
    return out, width


def _strip_headers_footers(pages_raw: list[list[dict]]) -> list[list[dict]]:
    """跨页剔除重复出现的页眉/页脚文本，以及页码类文本。"""
    zone_text_counter: Counter[str] = Counter()
    for page_raw in pages_raw:
        seen_this_page = set()
        for blk in page_raw:
            if blk["zone"] in ("top", "bottom") and blk["text"] not in seen_this_page:
                zone_text_counter[blk["text"]] += 1
                seen_this_page.add(blk["text"])
    repeated = {t for t, c in zone_text_counter.items() if c >= 2}

    result = []
    for page_raw in pages_raw:
        kept = [
            blk
            for blk in page_raw
            if not (blk["zone"] in ("top", "bottom") and (blk["text"] in repeated or _NUMERIC_ZONE_RE.match(blk["text"])))
        ]
        result.append(kept)
    return result


def _detect_column_split(blocks: list[dict], width: float) -> float | None:
    """判断是否存在有效分栏。

    先找候选分栏空白带，再要求两侧文本字符数都达到总字符数一定比例才采信
    ——真正的双栏内容大致对半分布，孤立块造成的伪分栏线两侧字符数会严重
    失衡。不按块宽度筛"宽块"：同一版面解析模型对单栏/双栏文本的分块粒度
    本身就不稳定，窄块在两种情况下都很常见，按宽度过滤并不可靠。
    """
    extents = [(b["bbox"][0], b["bbox"][2]) for b in blocks]
    split_x = _find_extent_gap(extents, width, config.COLUMN_GAP_BAND, config.COLUMN_MIN_GAP_WIDTH_RATIO)
    if split_x is None:
        return None
    total_chars = sum(len(b["text"]) for b in blocks)
    if total_chars == 0:
        return None
    left_chars = sum(len(b["text"]) for b in blocks if (b["bbox"][0] + b["bbox"][2]) / 2 < split_x)
    left_ratio = left_chars / total_chars
    if not (config.COLUMN_BALANCE_MIN_RATIO <= left_ratio <= 1 - config.COLUMN_BALANCE_MIN_RATIO):
        return None
    return split_x


def _classify_native_block_type(avg_size: float, body_size: float, zone: str) -> str:
    if body_size and avg_size > body_size * 1.15:
        return "heading"
    if zone == "bottom" and body_size and avg_size < body_size * 0.9:
        return "footnote"
    return "paragraph"


def _order_native_page(raw_blocks: list[dict], width: float, force_layout: str) -> tuple[list[dict], str]:
    if force_layout == "single":
        mode = "single"
    elif force_layout == "double":
        mode = "double"
    else:
        split_x = _detect_column_split(raw_blocks, width)
        mode = "double" if split_x is not None else "single"

    if mode == "single":
        ordered = sorted(raw_blocks, key=lambda b: b["bbox"][1])
        for b in ordered:
            b["column"] = None
        return ordered, mode

    split_x = _detect_column_split(raw_blocks, width)
    if split_x is None:
        split_x = width / 2

    col_width_est = width / 2
    spanning, left, right = [], [], []
    for b in raw_blocks:
        x0, _, x1, _ = b["bbox"]
        if (x1 - x0) > col_width_est * 1.4:
            spanning.append(b)
        else:
            center = (x0 + x1) / 2
            if center < split_x:
                b["column"] = "left"
                left.append(b)
            else:
                b["column"] = "right"
                right.append(b)

    left.sort(key=lambda b: b["bbox"][1])
    right.sort(key=lambda b: b["bbox"][1])
    spanning.sort(key=lambda b: b["bbox"][1])
    for b in spanning:
        b["column"] = "span"

    if not spanning:
        return left + right, mode

    ordered = []
    li = ri = 0
    for sp in spanning:
        sp_y = sp["bbox"][1]
        while li < len(left) and left[li]["bbox"][1] < sp_y:
            ordered.append(left[li])
            li += 1
        while ri < len(right) and right[ri]["bbox"][1] < sp_y:
            ordered.append(right[ri])
            ri += 1
        ordered.append(sp)
    ordered.extend(left[li:])
    ordered.extend(right[ri:])
    return ordered, mode


def _finalize_native_page(raw_blocks: list[dict], width: float, page_no: int, force_layout: str) -> tuple[list[dict], str]:
    if not raw_blocks:
        return [], "single"
    sizes = [b["avg_size"] for b in raw_blocks if b["avg_size"]]
    body_size = statistics.median(sizes) if sizes else 12.0
    for b in raw_blocks:
        b["block_type"] = _classify_native_block_type(b["avg_size"], body_size, b["zone"])

    ordered, mode = _order_native_page(raw_blocks, width, force_layout)
    finalized = [
        {
            "text": b["text"],
            "block_type": b["block_type"],
            "source_location": _source_location(page_no, mode, b.get("column")),
            "confidence": None,
        }
        for b in ordered
    ]
    return finalized, mode


# ---------------------------------------------------------------------------
# B类：无文字层 PDF（PaddleOCR 文本检测识别 + 独立版面检测 + XY-Cut阅读顺序）
# ---------------------------------------------------------------------------
#
# 说明：最初方案是直接调用 PPStructureV3 一体化管线，但在纯 CPU 环境
# （paddlepaddle 3.3.1 CPU 版）下该管线的并行调度会段错误崩溃（exit code 139）。
# 逐一排查确认：独立的 LayoutDetection（版面检测）与 PaddleOCR（文本检测+识别）
# 两个模型单独调用都完全正常。因此改为自行组合这两个模型：
#   1. LayoutDetection 得到区域框+标签（沿用与 PPStructureV3 相同的版面模型/标签体系）
#   2. PaddleOCR 得到文字行框+文本+置信度
#   3. 把文字行按中心点归属到对应版面区域，拼成整块文本、平均置信度
#   4. 用 paddlex 内部真正的 XY-Cut 算法（sort_by_xycut，纯 numpy 函数，
#      不经过模型执行器，没有崩溃风险）还原阅读顺序，双栏场景已内置正确处理
#
# 设备与 MKL-DNN：GPU 上推理快一到两个数量级，且不走 CPU 那条有已知算子兼容
# 问题的 MKL-DNN+PIR 执行路径。因此优先用 GPU；仅在回退到 CPU 时才需要禁用
# MKL-DNN 来规避 ConvertPirAttribute2RuntimeAttribute 崩溃。

_LAYOUT_PIPELINE = None
_OCR_PIPELINE = None
_PADDLE_DEVICE = None


def _resolve_device() -> str:
    """确定 paddle 推理设备（'gpu'/'cpu'），只解析一次并缓存。"""
    global _PADDLE_DEVICE
    if _PADDLE_DEVICE is None:
        setting = config.PADDLE_DEVICE
        if setting == "auto":
            try:
                import paddle

                has_gpu = paddle.device.is_compiled_with_cuda() and paddle.device.cuda.device_count() > 0
                _PADDLE_DEVICE = "gpu" if has_gpu else "cpu"
            except Exception:
                _PADDLE_DEVICE = "cpu"
        else:
            _PADDLE_DEVICE = setting
    return _PADDLE_DEVICE


def _model_common_kwargs() -> dict:
    """版面检测/OCR 模型共用的设备相关初始化参数。"""
    device = _resolve_device()
    kwargs: dict = {"device": device}
    if device == "cpu":
        # 仅 CPU 需要禁用 MKL-DNN 以规避 paddlepaddle 的已知算子兼容崩溃。
        kwargs["enable_mkldnn"] = False
    return kwargs


def _get_layout_pipeline():
    """惰性加载版面检测模型，避免非OCR场景（native/docx）也承担模型加载开销。"""
    global _LAYOUT_PIPELINE
    if _LAYOUT_PIPELINE is None:
        from paddleocr import LayoutDetection

        _LAYOUT_PIPELINE = LayoutDetection(**_model_common_kwargs())
    return _LAYOUT_PIPELINE


def _get_ocr_pipeline():
    """惰性加载文本检测+识别模型。"""
    global _OCR_PIPELINE
    if _OCR_PIPELINE is None:
        from paddleocr import PaddleOCR

        _OCR_PIPELINE = PaddleOCR(
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            lang=config.PADDLEOCR_LANG,
            **_model_common_kwargs(),
        )
    return _OCR_PIPELINE


# paddlex 版面标签 -> 本模块 block_type 映射（见 paddlex.inference.pipelines.layout_parsing.setting.BLOCK_LABEL_MAP）
_HEADING_LABELS = {"doc_title", "paragraph_title", "abstract_title", "reference_title", "content_title"}
_PARAGRAPH_LABELS = {"text", "content", "abstract", "reference", "formula", "algorithm", "aside_text"}
_FOOTNOTE_LABELS = {"footnote"}
_TABLE_LABELS = {"table"}
_CAPTION_LABELS = {"table_title", "chart_title", "figure_title", "figure_table_chart_title"}
_DROP_LABELS = {
    "image", "figure", "chart", "flowchart", "seal",
    "header", "header_image", "footer", "footer_image", "number",
}


def _label_to_block_type(label: str) -> str | None:
    if label in _HEADING_LABELS:
        return "heading"
    if label in _PARAGRAPH_LABELS:
        return "paragraph"
    if label in _FOOTNOTE_LABELS:
        return "footnote"
    if label in _TABLE_LABELS:
        return "table"
    if label in _CAPTION_LABELS:
        return "figure_caption"
    if label in _DROP_LABELS:
        return None
    return "other"


def _find_containing_region(cx: float, cy: float, regions: list[dict]) -> int | None:
    """返回中心点落在其内部、面积最小（最精确）的区域下标。"""
    candidates = []
    for i, r in enumerate(regions):
        x1, y1, x2, y2 = r["bbox"]
        if x1 <= cx <= x2 and y1 <= cy <= y2:
            candidates.append(((x2 - x1) * (y2 - y1), i))
    if not candidates:
        return None
    candidates.sort(key=lambda t: t[0])
    return candidates[0][1]


def _find_nearest_region(cx: float, cy: float, regions: list[dict]) -> int:
    best_i, best_dist = 0, None
    for i, r in enumerate(regions):
        x1, y1, x2, y2 = r["bbox"]
        dx = max(x1 - cx, 0, cx - x2)
        dy = max(y1 - cy, 0, cy - y2)
        dist = dx * dx + dy * dy
        if best_dist is None or dist < best_dist:
            best_dist, best_i = dist, i
    return best_i


def _run_structure(img: Image.Image) -> tuple[list[dict], float | None]:
    from paddlex.inference.pipelines.layout_parsing.xycut_enhanced.xycuts import sort_by_xycut

    layout_pipeline = _get_layout_pipeline()
    ocr_pipeline = _get_ocr_pipeline()

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = str(Path(tmp) / "page.png")
        img.save(tmp_path)
        layout_result = list(layout_pipeline.predict(tmp_path))[0]
        ocr_result = list(ocr_pipeline.predict(tmp_path))[0]

    regions = [{"label": b["label"], "bbox": tuple(b["coordinate"]), "lines": []} for b in layout_result["boxes"]]
    if not regions:
        return [], None

    for text, score, box in zip(ocr_result["rec_texts"], ocr_result["rec_scores"], ocr_result["rec_boxes"]):
        if not text.strip():
            continue
        x1, y1, x2, y2 = box
        cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
        region_idx = _find_containing_region(cx, cy, regions)
        if region_idx is None:
            region_idx = _find_nearest_region(cx, cy, regions)
        regions[region_idx]["lines"].append({"text": text, "score": float(score), "bbox": box})

    order = sort_by_xycut([r["bbox"] for r in regions], direction="horizontal", min_gap=1)

    blocks = []
    scores = []
    for idx in order:
        region = regions[idx]
        block_type = _label_to_block_type(region["label"])
        lines = sorted(region["lines"], key=lambda l: (l["bbox"][1], l["bbox"][0]))
        text = "".join(l["text"] for l in lines).strip()
        if block_type is None or not text:
            continue
        conf = sum(l["score"] for l in lines) / len(lines) if lines else None
        if conf is not None:
            scores.append(conf)
        blocks.append({"text": text, "block_type": block_type, "bbox": region["bbox"], "confidence": conf})

    page_conf = sum(scores) / len(scores) if scores else None
    return blocks, page_conf


def _ocr_and_order(img: Image.Image, force_layout: str) -> tuple[list[dict], float | None, str]:
    raw_blocks, page_conf = _run_structure(img)

    filtered = []
    for b in raw_blocks:
        conf = b["confidence"]
        if conf is not None and conf < config.OCR_MIN_BLOCK_CONFIDENCE and len(b["text"]) < config.OCR_MIN_BLOCK_CHARS:
            continue  # 极低置信度+极短文本，视为乱码丢弃
        filtered.append(b)

    if force_layout in ("single", "double"):
        mode = force_layout
        split_x = None
    else:
        text_like = [b for b in filtered if b["block_type"] in ("paragraph", "heading")]
        split_x = _detect_column_split(text_like, img.width)
        mode = "double" if split_x is not None else "single"

    if mode == "double":
        if split_x is None:
            split_x = img.width / 2
        for b in filtered:
            cx = (b["bbox"][0] + b["bbox"][2]) / 2
            b["column"] = "left" if cx < split_x else "right"
    else:
        for b in filtered:
            b["column"] = None

    return filtered, page_conf, mode


# ---------------------------------------------------------------------------
# B类：跨页（横版两页拼接）检测与拆分
# ---------------------------------------------------------------------------

def _render_page_image(page: "fitz.Page", dpi: int) -> Image.Image:
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix, colorspace=fitz.csRGB)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def _find_pixel_gap(img: Image.Image, band: tuple[float, float]) -> tuple[float, float] | None:
    gray = np.asarray(img.convert("L"), dtype=np.float32)
    dark = gray < 200
    col_density = dark.mean(axis=0)
    w = col_density.shape[0]
    band_i0 = max(0, int(band[0] * w))
    band_i1 = min(w, int(band[1] * w))
    threshold = 0.01

    best: tuple[int, int] | None = None
    start: int | None = None
    for i in range(band_i0, band_i1):
        if col_density[i] < threshold:
            if start is None:
                start = i
        elif start is not None:
            if best is None or (i - start) > (best[1] - best[0]):
                best = (start, i)
            start = None
    if start is not None:
        if best is None or (band_i1 - start) > (best[1] - best[0]):
            best = (start, band_i1)

    if best is None:
        return None
    return best[0] / w, best[1] / w


def _maybe_split_spread(img: Image.Image, spread_order: str, warnings: list[str], phys_page_no: int) -> list[Image.Image]:
    aspect = img.width / img.height
    if aspect <= config.SPREAD_ASPECT_RATIO_THRESHOLD:
        return [img]

    gap = _find_pixel_gap(img, band=(0.3, 0.7))
    if gap is None or (gap[1] - gap[0]) < config.SPREAD_GAP_WIDTH_RATIO:
        return [img]

    split_x = int((gap[0] + gap[1]) / 2 * img.width)
    left = img.crop((0, 0, split_x, img.height))
    right = img.crop((split_x, 0, img.width, img.height))
    warnings.append(
        f"第{phys_page_no}页检测为横版两页拼接，已从中缝拆分为两个逻辑页（左前右后），"
        f"请人工确认页面顺序{'（含首页，可能为封面）' if phys_page_no == 1 or spread_order == 'cover_first' else ''}"
    )
    return [left, right]
