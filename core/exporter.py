"""Excel 导出模块。

把已落库（core/workflow/persist.py::persist_result 写入的 records/issues 两表）
的一轮校对结果导出为 Excel 清单：按页码排序、高优先级标红、引文类标浅黄。

数据一律从数据库读，不依赖内存 ClassifiedResult——"处理状态"列要反映用户在
Streamlit 界面上采纳/拒绝的最新决策，这个状态只存在库里（set_issue_status 写的）。
本模块不重建 record、不重写统计字段（persist_result 已经写好），只在导出后
回填 result_path。

`export_issues_to_excel(record_id: int, db_path=None) -> Path` 是唯一对外
入口。签名比占位多一个 db_path=None——对齐全仓库 db.models 函数"支持传临时
库测试"的惯例，本函数内部要调 get_issues/get_records/update_record_stats，
必须能透传测试用的临时库路径。

三条关键设计决策：

1. 数据一律从数据库读，不依赖内存 ClassifiedResult：用 get_issues(record_id)
   + get_records() 查询。这也让本函数自包含，能被历史记录页对任意历史
   record 复用。
2. 不重建 record、不重写统计字段，只回填 result_path：persist_result 已经把
   total_issues/count_confirmed 等全部统计字段写进 records 表了（issues.record_id
   NOT NULL 外键约束要求先有 record）。本函数生成文件后只调一次
   update_record_stats(record_id, result_path=str(path), db_path=db_path)。
   **result_path 只写不读**（没有任何代码按它去取回文件），所以 _prune_exports_dir
   删掉旧导出文件不影响任何已有记录，只是那些行的路径指向一个已不存在的文件。
3. 行顺序直接用 get_issues 的 issue_id 升序，不解析 page_location（"第N页"
   之类的 TEXT）重排——core/classifier/ 已经把 ClassifiedResult.issues 排好序
   （页码升序/同页按block_index/未定位排最后），persist_result 按该顺序逐条
   add_issue 写入，issue_id 自增顺序本身就是正确阅读顺序。

高优先级标红 与 引文类标浅黄 的冲突取舍：优先级覆盖规则会真的造出"引文类+
高优先级"的条目（issue_type 命中政治敏感性表述/民族与地名规范时，无论落在
哪层 priority 强制改为'高'，哪怕是引文类）。这种情况按标红处理（红色是更强
的警示信号），本文件里用 if/elif 保证互斥（红色判断在前）。
tests/test_exporter.py::test_export_quotation_and_high_priority_conflict_resolves_to_red
专门构造这种条目钉住这个取舍，避免以后改样式时不小心变成随机结果。

样式细节：高优先级整行填 FFFFC7CE（浅红），引文类整行填 FFFFF2CC（浅黄），
表头加粗+浅灰底+freeze_panes；"原文"/"修改建议"两列加宽+自动换行。颜色等
样式常量硬编码在本文件，没有进 config.py（非必要不加可调项）。0条issue的
record仍生成合法的（只有表头的）xlsx，不抛异常；不存在的 record_id 抛
ValueError。

批注与导出：

1. "拒绝理由"是通用"批注"，与采纳/拒绝状态解耦——详见
   core/workflow/CLAUDE.md"批注与状态解耦"一节。issues 表列名为 note
   （db/database.py::init_db() 里的 _migrate_reject_reason_to_note() 负责
   把历史库的旧列名 reject_reason 幂等迁移过来）。
2. Excel导出只导出"已采纳"的问题：界面上通常只勾选"采纳"，"待处理"/"已拒绝"
   不该出现在导出结果里，export_issues_to_excel 内部按 status == "已采纳"
   过滤 get_issues() 的结果，无条件过滤，不做成可选参数（没有"导出全部状态"
   的需求，加一个用不到的开关只会增加复杂度）。"处理状态"列直接显示原始
   status（导出结果里恒为"已采纳"），独立的"批注"列显示 note（未写则留空）。

"页码/位置"列（issues.page_location）一律是PDF物理页码（"第N页左栏"），是导出
清单里唯一的定位列，也是校对时对着PDF翻页找问题的唯一坐标。_HEADER 共8列。

回归测试：tests/test_exporter.py::test_export_note_column_shows_value_and_status_is_plain、
test_export_only_includes_accepted_issues、
test_export_record_with_no_accepted_issues_generates_header_only_file。
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config
from db.models import get_issues, get_records, update_record_stats

_HEADER = ("页码/位置", "原文", "问题类型", "优先级", "分层标注", "修改建议", "处理状态", "批注")

_HEADER_FILL = PatternFill("solid", fgColor="FFD9D9D9")
_HIGH_PRIORITY_FILL = PatternFill("solid", fgColor="FFFFC7CE")
_QUOTATION_FILL = PatternFill("solid", fgColor="FFFFF2CC")

_ILLEGAL_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')

_EXPORTED_STATUS = "已采纳"


def _sanitize_filename(name: str) -> str:
    return _ILLEGAL_FILENAME_CHARS.sub("_", name)


def _prune_exports_dir() -> None:
    """只保留 config.EXPORTS_DIR 下最近修改的 config.EXPORTS_RETENTION_COUNT 个文件。

    与上传目录同一套策略（ui/actions.py::prune_uploads_dir），理由也一样：导出的
    Excel 生成后立刻由下载按钮交给用户，服务端这份只是中转，没有任何代码会再读回来
    （records.result_path 只写不读），不删会无限堆积。放在本模块而不是 ui/ 里，是因为
    三个页面的导出都汇到这一个函数，挂在这里才不会漏掉某个入口。
    """
    files = [f for f in config.EXPORTS_DIR.iterdir() if f.is_file()]
    if len(files) <= config.EXPORTS_RETENTION_COUNT:
        return
    files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    for stale in files[config.EXPORTS_RETENTION_COUNT:]:
        stale.unlink(missing_ok=True)


def export_issues_to_excel(record_id: int, db_path=None) -> Path:
    """将指定流程记录下已采纳的问题导出为 Excel 文件，回填 result_path，返回文件路径。

    只导出状态为"已采纳"的问题：真实使用中发现，每次都是先在界面上只勾选采纳、
    其余（待处理/已拒绝）导出后手动从Excel里删掉，导出时直接过滤更省事。
    """
    records = get_records(db_path=db_path)
    record = next((r for r in records if r["record_id"] == record_id), None)
    if record is None:
        raise ValueError(f"record_id {record_id} 不存在")

    issues = [i for i in get_issues(record_id, db_path=db_path) if i["status"] == _EXPORTED_STATUS]

    wb = Workbook()
    ws = wb.active
    ws.title = "校对清单"

    ws.append(_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    ws.freeze_panes = "A2"

    for row in issues:
        ws.append(
            (
                row["page_location"] or "未定位",
                row["original_text"],
                row["issue_type"],
                row["priority"],
                row["layer"],
                row["suggestion"],
                row["status"],
                row["note"] or "",
            )
        )
        excel_row = ws.max_row
        # 高优先级标红优先于引文类标浅黄：优先级覆盖规则会真的造出
        # "引文类+高优先级"的条目（政治敏感/民族地名类），红色是更强的警示信号。
        if row["priority"] == config.PRIORITY_HIGH:
            fill = _HIGH_PRIORITY_FILL
        elif row["layer"] == config.LAYER_QUOTATION:
            fill = _QUOTATION_FILL
        else:
            fill = None
        if fill is not None:
            for cell in ws[excel_row]:
                cell.fill = fill

    wrap = Alignment(wrap_text=True, vertical="top")
    for col_idx, width in zip(range(1, 9), (14, 50, 16, 10, 12, 50, 12, 30)):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = width
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = wrap

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    doc_name = _sanitize_filename(record["doc_name"] or "未命名文档")
    filename = f"{timestamp}_{doc_name}_校对清单.xlsx"
    path = config.EXPORTS_DIR / filename
    wb.save(path)
    _prune_exports_dir()  # 保留刚生成的这份（按修改时间它最新），只清更早的

    update_record_stats(record_id, result_path=str(path), db_path=db_path)
    return path
