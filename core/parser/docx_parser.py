"""C类：Word 文档解析（从 core/parser.py 拆分而来，逻辑未改动）。"""

from __future__ import annotations

import io
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import docx

from core.parser._types import ParsedBlock, ParsedDocument

_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", _RELS_NS)  # 避免重新序列化时ElementTree生成丑陋的ns0:前缀


def _resolve_relationship_target(rels_member_name: str, target: str) -> str:
    """把 .rels 文件里的 Target（相对路径）解析成 zip 内的成员名（不带开头斜杠）。

    OPC规范：关系记录的Target是相对"拥有这份.rels文件的part所在目录"解析的，比如
    word/_rels/document.xml.rels 里的Target是相对 word/ 目录解析的。用posixpath纯
    字符串处理（不碰文件系统），../这类相对路径能正确归一化。
    """
    if target.startswith("/"):
        return target.lstrip("/")
    base_dir = posixpath.dirname(posixpath.dirname(rels_member_name))  # .../_rels/x.rels -> ...
    return posixpath.normpath(posixpath.join(base_dir, target))


def _repair_dangling_relationships(path: Path) -> Path | io.BytesIO:
    """摘掉docx内部各 .rels 文件里指向"zip中实际不存在的部件"的断链关系记录。

    起因（真实文档诊断）：python-docx 打开文件时会急切读取所有关系记录指向的部件，
    一条断链记录（常见于PDF转Word等转换工具生成的docx——比如某张图片的关系Target
    被写成了字面意义上的"NULL"）会导致 KeyError、整个文件都打不开。Word本身对这种
    情况很宽容：那张图就是不显示，正文照常打开，不会弹出明显的损坏提示。这里在原始
    zip层面直接移除断链的关系记录，模拟Word的这种容错——docx_parser.parse_docx 只读
    取段落文字，不涉及图片/媒体部件，摘掉这些记录对解析结果没有任何影响。

    TargetMode="External"（如指向网址的超链接）的关系记录本来就不该在zip里找到对应
    部件，不算断链，跳过不处理。

    没有任何断链时原样返回 path，不产生任何多余的zip重写开销；有断链时返回内存里
    修复好的字节流（io.BytesIO），docx.Document() 能直接接收文件路径或类文件对象，
    不需要另外写临时文件到磁盘。
    """
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        rels_members = [n for n in names if n.endswith(".rels")]

        patched: dict[str, bytes] = {}
        for rels_name in rels_members:
            root = ET.fromstring(zf.read(rels_name))
            changed = False
            for rel in list(root):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target")
                if target is None:
                    continue
                if _resolve_relationship_target(rels_name, target) not in names:
                    root.remove(rel)
                    changed = True
            if changed:
                patched[rels_name] = ET.tostring(root, encoding="UTF-8", xml_declaration=True)

        if not patched:
            return path  # 没有断链，不需要修复

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as out:
            for item in zf.infolist():
                data = patched.get(item.filename, zf.read(item.filename))
                out.writestr(item, data)
        buffer.seek(0)
        return buffer


def parse_docx(path: Path) -> ParsedDocument:
    """Word 解析：直接按段落顺序读取，不涉及分栏/OCR，是三条通道里最简单的一条。

    先按正常路径打开；只有真的撞上"断链关系记录导致KeyError"这种具体场景才退化到
    _repair_dangling_relationships 修复后重试——不在每次调用时都提前扫描.rels（多数
    文档没有这个问题，没必要为小概率场景增加每次解析的开销）。修复仍然失败（比如
    根本不是这个原因导致的KeyError）就让异常照常往上抛，走 core/parser/__init__.py
    既有的异常处理路径，行为不比修复前更差。
    """
    try:
        document = docx.Document(str(path))
    except KeyError:
        document = docx.Document(_repair_dangling_relationships(path))
    blocks: list[ParsedBlock] = []
    idx = 0  # 剔除空段落后的顺序号，与下面循环里的 i（原始段落序号）不是同一个数
    for i, para in enumerate(document.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue  # 空段落（纯换行/纯空格）直接跳过，不产出 ParsedBlock
        # para.style 是 Word 里作者手动设置的段落样式（点击"标题1"/"正文"等按钮时写入的元数据），
        # 不是内容分析结果。这里只做字符串匹配：样式名以"Heading"开头或等于"Title"就算标题，
        # 否则一律归为正文——如果作者没用 Word 内置标题样式，这里识别不出来。
        style_name = para.style.name if para.style is not None else ""
        block_type = "heading" if (style_name.startswith("Heading") or style_name == "Title") else "paragraph"
        blocks.append(
            ParsedBlock(
                page=0,               # Word 无页码概念，统一填0
                block_index=idx,
                text=text,
                block_type=block_type,
                source_location=f"第{i}段",  # 用原始段落号定位，方便用户在Word里核对
                ocr_confidence=None,   # Word 文本非OCR来源
            )
        )
        idx += 1
    return ParsedDocument(
        file_name=path.name,
        file_type="docx",
        total_pages=len(document.paragraphs),  # 用段落总数代替页数
        blocks=blocks,
        layout_mode="single",   # Word 文档不做分栏检测，固定单栏
        text_source="native",   # 固定为原生文本（非OCR）
        warnings=[],
    )
