"""阶段8验收测试：Excel 导出（core/exporter.py）。

全部不联网不耗API额度：core/exporter.py 直接对临时数据库+临时导出目录操作，
用 openpyxl.load_workbook 打开生成文件做精确断言；app.py 的导出按钮用
streamlit.testing.v1.AppTest 冒烟，打桩掉 core.exporter.export_issues_to_excel。
"""

import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from core.exporter import export_issues_to_excel
from core.parser import ParsedDocument
from db import database
from db.models import add_issue, create_record, create_task, get_records, get_tasks

_HEADER = ("页码/位置", "文档页码", "原文", "问题类型", "优先级", "分层标注", "修改建议", "处理状态", "批注")


@pytest.fixture
def db_path(tmp_path):
    path = tmp_path / "test_app.db"
    database.init_db(path)
    return path


def _enter_task(db_path) -> int:
    """预置一个当前任务，跳过 app.py 的任务闸门。

    app.py 是"先选任务再选功能"的线性流程——没有当前任务时只渲染任务选择界面、
    连"功能入口"radio 都不存在。这些UI冒烟测试要测的是任务之内的各个页面，所以直接
    预置 session_state["task_id"]；闸门本身由 tests/test_app_tasks.py 单独覆盖。
    """
    tasks = get_tasks(db_path=db_path)
    return tasks[0]["task_id"] if tasks else create_task("测试任务", db_path=db_path)


def _task(db_path) -> int:
    """建一个任务——records.task_id 是必填的，任何 create_record 之前都要先有任务。"""
    return create_task("测试任务", db_path=db_path)


def _build_record_with_issues(db_path):
    """构造覆盖四层、高优先级、引文类、引文+高优先级、未定位的记录。

    导出只导出"已采纳"的问题（真实使用中发现的行为变更，见 core/exporter.py），
    所以这里全部用"已采纳"状态构造，以便header/顺序/颜色等断言仍能覆盖到这些行；
    "只导出已采纳"这条行为本身由 test_export_only_includes_accepted_issues 单独覆盖。
    """
    record_id = create_record(
        task_id=_task(db_path), doc_name="测试文档.pdf", doc_version="v1",
        task_type="标准校对", db_path=db_path,
    )

    issue_specs = [
        dict(
            original_text="确定性问题原文",
            issue_type="错别字与拼写",
            priority=config.PRIORITY_MEDIUM,
            layer=config.LAYER_CONFIRMED,
            suggestion="改为正确写法",
            page_location="第1页",
            status="已采纳",
        ),
        dict(
            original_text="存疑问题原文",
            issue_type="常识与事实性错误",
            priority=config.PRIORITY_MEDIUM,
            layer=config.LAYER_DOUBTFUL,
            suggestion="存疑，建议人工核实：可能有误",
            page_location="第2页",
            status="已采纳",
        ),
        dict(
            original_text="引文问题原文",
            issue_type="引用与成语准确性",
            priority=config.PRIORITY_LOW,
            layer=config.LAYER_QUOTATION,
            suggestion="原文照录，不建议改动。",
            page_location="第3页",
            status="已采纳",
        ),
        dict(
            original_text="风格问题原文",
            issue_type="表达润色",
            priority=config.PRIORITY_OPTIONAL,
            layer=config.LAYER_OPTIONAL,
            suggestion="建议润色，更通顺",
            page_location="第4页",
            status="已采纳",
        ),
        dict(
            original_text="高优先级确定性问题原文",
            issue_type="政治敏感性表述",
            priority=config.PRIORITY_HIGH,
            layer=config.LAYER_CONFIRMED,
            suggestion="需修改的表述",
            page_location="第5页",
            status="已采纳",
            note="已核实，确实需要修改",
        ),
        dict(
            original_text="引文且高优先级问题原文",
            issue_type="民族与地名规范",
            priority=config.PRIORITY_HIGH,
            layer=config.LAYER_QUOTATION,
            suggestion="原文照录，不建议改动。",
            page_location=None,
            status="已采纳",
        ),
    ]

    issue_ids = []
    for spec in issue_specs:
        issue_id = add_issue(
            record_id=record_id,
            page_location=spec["page_location"],
            original_text=spec["original_text"],
            issue_type=spec["issue_type"],
            priority=spec["priority"],
            layer=spec["layer"],
            suggestion=spec["suggestion"],
            status=spec["status"],
            note=spec.get("note"),
            db_path=db_path,
        )
        issue_ids.append(issue_id)

    return record_id, issue_specs, issue_ids


# ---------------------------------------------------------------------------
# core/exporter.py 单元测试
# ---------------------------------------------------------------------------

def test_export_creates_workbook_with_header_and_ordered_rows(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, issue_ids = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)

    assert path.exists()
    assert path.parent == tmp_path

    wb = load_workbook(path)
    ws = wb.active

    header = tuple(cell.value for cell in ws[1])
    assert header == _HEADER

    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == len(issue_specs)
    for row, spec in zip(data_rows, issue_specs):
        assert row[2] == spec["original_text"]  # 顺序须与 issue_id(插入顺序)一致


def test_export_page_location_none_shows_unlocated(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    row = next(
        r for r in ws.iter_rows(min_row=2, values_only=True)
        if r[2] == "引文且高优先级问题原文"
    )
    assert row[0] == "未定位"


def test_export_doc_page_column_blank_when_not_extracted(db_path, tmp_path, monkeypatch):
    """_build_record_with_issues 构造的记录全部没传 doc_page，导出应留空——
    不该用PDF页码顶替（PDF页码信息已经在"页码/位置"列里）。"""
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    for row in ws.iter_rows(min_row=2, values_only=True):
        assert row[1] is None or row[1] == ""


def test_export_doc_page_column_shows_value_when_extracted(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id = create_record(task_id=_task(db_path), doc_name="期刊.pdf", doc_version="", task_type="标准校对", db_path=db_path)
    add_issue(
        record_id=record_id, page_location="文档第12页左栏", doc_page="12", original_text="期刊正文",
        issue_type="错别字与拼写", priority=config.PRIORITY_MEDIUM, layer=config.LAYER_CONFIRMED,
        suggestion="改为正确写法", status="已采纳", db_path=db_path,
    )

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[2] == "期刊正文")
    assert row[0] == "文档第12页左栏"
    assert row[1] == "12"


def test_export_note_column_shows_value_and_status_is_plain(db_path, tmp_path, monkeypatch):
    """批注不再和拒绝绑定：处理状态列只显示原始status，批注单独一列，无批注时留空。"""
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    noted_row = next(
        r for r in ws.iter_rows(min_row=2, values_only=True)
        if r[2] == "高优先级确定性问题原文"
    )
    assert noted_row[7] == "已采纳"  # 处理状态不再拼接批注内容
    assert noted_row[8] == "已核实，确实需要修改"

    plain_row = next(
        r for r in ws.iter_rows(min_row=2, values_only=True)
        if r[2] == "存疑问题原文"
    )
    assert plain_row[7] == "已采纳"
    assert not plain_row[8]  # 未写批注，写入的是空字符串；openpyxl读回后为None，语义上等同"留空"


def test_export_only_includes_accepted_issues(db_path, tmp_path, monkeypatch):
    """真实使用中发现的行为变更：待处理/已拒绝的问题不应出现在导出结果里。"""
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id = create_record(task_id=_task(db_path), doc_name="混合状态文档.pdf", doc_version="", task_type="标准校对", db_path=db_path)
    add_issue(
        record_id=record_id, page_location="第1页", original_text="待处理问题",
        issue_type="错别字与拼写", priority=config.PRIORITY_MEDIUM, layer=config.LAYER_CONFIRMED,
        suggestion="改为正确写法", status="待处理", db_path=db_path,
    )
    add_issue(
        record_id=record_id, page_location="第2页", original_text="已拒绝问题",
        issue_type="错别字与拼写", priority=config.PRIORITY_MEDIUM, layer=config.LAYER_CONFIRMED,
        suggestion="改为正确写法", status="已拒绝", note="确认原文没错", db_path=db_path,
    )
    add_issue(
        record_id=record_id, page_location="第3页", original_text="已采纳问题",
        issue_type="错别字与拼写", priority=config.PRIORITY_MEDIUM, layer=config.LAYER_CONFIRMED,
        suggestion="改为正确写法", status="已采纳", db_path=db_path,
    )

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 1
    assert data_rows[0][2] == "已采纳问题"


def test_export_record_with_no_accepted_issues_generates_header_only_file(db_path, tmp_path, monkeypatch):
    """记录里有问题，但一条都没被采纳——导出应是只有表头的空清单，不抛异常。"""
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id = create_record(task_id=_task(db_path), doc_name="全部待处理文档.pdf", doc_version="", task_type="标准校对", db_path=db_path)
    add_issue(
        record_id=record_id, page_location="第1页", original_text="待处理问题",
        issue_type="错别字与拼写", priority=config.PRIORITY_MEDIUM, layer=config.LAYER_CONFIRMED,
        suggestion="改为正确写法", status="待处理", db_path=db_path,
    )

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    assert list(ws.iter_rows(min_row=2, values_only=True)) == []


def test_export_high_priority_rows_are_red(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    def _row_by_text(text):
        for r in ws.iter_rows(min_row=2):
            if r[2].value == text:
                return r
        raise AssertionError(f"未找到原文: {text}")

    high_priority_row = _row_by_text("高优先级确定性问题原文")
    for cell in high_priority_row:
        assert cell.fill.fgColor.rgb == "FFFFC7CE"


def test_export_quotation_only_rows_are_yellow(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    def _row_by_text(text):
        for r in ws.iter_rows(min_row=2):
            if r[2].value == text:
                return r
        raise AssertionError(f"未找到原文: {text}")

    quotation_row = _row_by_text("引文问题原文")
    for cell in quotation_row:
        assert cell.fill.fgColor.rgb == "FFFFF2CC"


def test_export_quotation_and_high_priority_conflict_resolves_to_red(db_path, tmp_path, monkeypatch):
    """决策5：高优先级标红 优先于 引文类标浅黄。"""
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    def _row_by_text(text):
        for r in ws.iter_rows(min_row=2):
            if r[2].value == text:
                return r
        raise AssertionError(f"未找到原文: {text}")

    conflict_row = _row_by_text("引文且高优先级问题原文")
    for cell in conflict_row:
        assert cell.fill.fgColor.rgb == "FFFFC7CE"


def test_export_normal_rows_have_no_fill(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)
    ws = load_workbook(path).active

    def _row_by_text(text):
        for r in ws.iter_rows(min_row=2):
            if r[2].value == text:
                return r
        raise AssertionError(f"未找到原文: {text}")

    normal_row = _row_by_text("确定性问题原文")
    for cell in normal_row:
        assert cell.fill.fgColor.rgb == "00000000"


def test_export_backfills_result_path_on_record(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id, issue_specs, _ = _build_record_with_issues(db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)

    records = {r["record_id"]: r for r in get_records(db_path=db_path)}
    assert records[record_id]["result_path"] == str(path)


# ---------------------------------------------------------------------------
# 边界情况
# ---------------------------------------------------------------------------

def test_export_record_with_zero_issues_still_generates_header_only_file(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)
    record_id = create_record(task_id=_task(db_path), doc_name="空文档.pdf", doc_version="", task_type="标准校对", db_path=db_path)

    path = export_issues_to_excel(record_id, db_path=db_path)

    assert path.exists()
    ws = load_workbook(path).active
    header = tuple(cell.value for cell in ws[1])
    assert header == _HEADER
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert data_rows == []

    records = {r["record_id"]: r for r in get_records(db_path=db_path)}
    assert records[record_id]["result_path"] == str(path)


def test_export_nonexistent_record_id_raises(db_path, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path)

    with pytest.raises(ValueError):
        export_issues_to_excel(999999, db_path=db_path)


# ---------------------------------------------------------------------------
# app.py UI 冒烟
# ---------------------------------------------------------------------------

def test_app_export_button_downloads_file(tmp_path, monkeypatch, db_path):
    monkeypatch.setattr(config, "DB_PATH", db_path)
    from unittest.mock import MagicMock

    from core.classifier import ClassifiedIssue, ClassifiedResult
    from core.parser import ParsedDocument
    from streamlit.testing.v1 import AppTest

    fake_issue = ClassifiedIssue(
        original_text="示例原文",
        issue_type="标点符号问题",
        suggestion="建议修改",
        reason="示例依据",
        block_index=0,
        page_location="第1页",
        chunk_index=0,
        located=True,
        layer=config.LAYER_CONFIRMED,
        priority=config.PRIORITY_MEDIUM,
        layer_notes=["测试用例构造"],
        llm_category="normal",
        llm_confidence="high",
        original_suggestion="建议修改",
    )
    fake_result = ClassifiedResult(
        issues=[fake_issue],
        stats={
            "total_issues": 1,
            "count_confirmed": 1,
            "count_doubtful": 0,
            "count_quotation": 0,
            "count_optional": 0,
            "high_priority_count": 0,
        },
        warnings=[],
    )

    export_dir = tmp_path / "exports"
    export_dir.mkdir()
    fake_export_path = export_dir / "fake_export.xlsx"
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(_HEADER)
    wb.save(fake_export_path)

    uploads_dir = tmp_path / "uploads"
    uploads_dir.mkdir()
    monkeypatch.setattr(config, "UPLOADS_DIR", uploads_dir)

    fake_parsed = MagicMock(spec=ParsedDocument)

    with patch("core.workflow.run_standard_proofread", return_value=(fake_result, fake_parsed)), \
         patch("core.workflow.persist_result", return_value=(1, [101])), \
         patch("core.followup.get_followup_history", return_value=[]), \
         patch("core.feedback_rules.regenerate_rejection_rules"), \
         patch("core.exporter.export_issues_to_excel", return_value=fake_export_path) as mock_export:
        at = AppTest.from_file(str(Path(__file__).resolve().parent.parent / "app.py"))
        at.session_state["task_id"] = _enter_task(db_path)
        at.run()

        at.file_uploader[0].upload("test.pdf", b"dummy pdf bytes", "application/pdf").run()
        start_button = next(b for b in at.button if b.label == "开始校对")
        start_button.click().run()
        assert not at.exception

        # 标准校对页导出成功后会触发 regenerate_rejection_rules（已打桩），避免真发起LLM调用。
        export_button = next(b for b in at.button if b.label == "导出Excel")
        export_button.click().run()

        # AppTest（当前版本）不支持 st.download_button 元素访问（无 at.download_button
        # 属性），退化为"不崩溃 + 导出函数被正确调用一次 + 成功提示展示"这条较弱断言。
        assert not at.exception
        mock_export.assert_called_once_with(1)
        assert len(at.success) >= 1
